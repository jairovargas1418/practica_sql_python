import pyodbc
import pandas as pd
import matplotlib.pyplot as plt

from datetime import datetime

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.drawing.image import Image


# =====================================================
# CONEXION SQL SERVER
# =====================================================

conexion = pyodbc.connect(
    "DRIVER={ODBC Driver 17 for SQL Server};"
    "SERVER=PC-580;"
    "DATABASE=BD_CLINICA_PRACTICA;"
    "Trusted_Connection=yes;"
)

print("✅ Conexión exitosa")


# =====================================================
# PACIENTES POR CIUDAD
# =====================================================

sql_pacientes_ciudad = """
SELECT
    Ciudad,
    COUNT(*) AS TotalPacientes
FROM Paciente
GROUP BY Ciudad
ORDER BY TotalPacientes DESC
"""

df_pacientes_ciudad = pd.read_sql(
    sql_pacientes_ciudad,
    conexion
)


# =====================================================
# CITAS POR MEDICO
# =====================================================

sql_citas_medico = """
SELECT
    m.NombreCompleto AS Medico,
    COUNT(*) AS TotalCitas
FROM Cita c
INNER JOIN Medico m
    ON c.IdMedico = m.IdMedico
GROUP BY m.NombreCompleto
ORDER BY TotalCitas DESC
"""

df_citas_medico = pd.read_sql(
    sql_citas_medico,
    conexion
)
df_ranking_medicos = df_citas_medico.sort_values(
    by="TotalCitas",
    ascending=False
)

# =====================================================
# CITAS POR ESPECIALIDAD
# =====================================================

sql_especialidad = """
SELECT
    e.Nombre AS Especialidad,
    COUNT(*) AS TotalCitas
FROM Cita c
INNER JOIN Medico m
    ON c.IdMedico = m.IdMedico
INNER JOIN Especialidad e
    ON m.IdEspecialidad = e.IdEspecialidad
GROUP BY e.Nombre
ORDER BY TotalCitas DESC
"""

df_especialidad = pd.read_sql(
    sql_especialidad,
    conexion
)
df_ranking_especialidades = (
    df_especialidad.sort_values(
        by="TotalCitas",
        ascending=False
    )
)

# =====================================================
# FACTURACION
# =====================================================

sql_facturacion = """
SELECT
    SUM(Valor) AS FacturacionTotal
FROM Factura
"""

df_facturacion = pd.read_sql(
    sql_facturacion,
    conexion
)


# =====================================================
# RESUMEN GENERAL
# =====================================================

sql_resumen = """
SELECT
    (SELECT COUNT(*) FROM Paciente) AS TotalPacientes,
    (SELECT COUNT(*) FROM Medico) AS TotalMedicos,
    (SELECT COUNT(*) FROM Cita) AS TotalCitas,
    (SELECT COUNT(*) FROM Consulta) AS TotalConsultas,
    (SELECT COUNT(*) FROM Factura) AS TotalFacturas
"""

df_resumen = pd.read_sql(
    sql_resumen,
    conexion
)


# =====================================================
# KPIs
# =====================================================

total_pacientes = int(df_resumen.loc[0, 'TotalPacientes'])
total_medicos = int(df_resumen.loc[0, 'TotalMedicos'])
total_citas = int(df_resumen.loc[0, 'TotalCitas'])

facturacion = float(
    df_facturacion.loc[0, 'FacturacionTotal']
)

df_kpi = pd.DataFrame(
    {
        "Indicador": [
            "Pacientes",
            "Medicos",
            "Citas",
            "Facturacion"
        ],
        "Valor": [
            total_pacientes,
            total_medicos,
            total_citas,
            facturacion
        ]
    }
)


# =====================================================
# GRAFICO PACIENTES POR CIUDAD
# =====================================================

plt.figure(figsize=(8, 5))

plt.bar(
    df_pacientes_ciudad["Ciudad"],
    df_pacientes_ciudad["TotalPacientes"]
)

plt.title("Pacientes por Ciudad")
plt.xlabel("Ciudad")
plt.ylabel("Cantidad")

plt.tight_layout()

plt.savefig(
    "Reportes/pacientes_por_ciudad.png"
)

plt.close()


# =====================================================
# GRAFICO ESPECIALIDADES
# =====================================================

plt.figure(figsize=(8, 5))

plt.bar(
    df_especialidad["Especialidad"],
    df_especialidad["TotalCitas"]
)

plt.title("Citas por Especialidad")
plt.xlabel("Especialidad")
plt.ylabel("Total Citas")

plt.xticks(rotation=30)

plt.tight_layout()

plt.savefig(
    "Reportes/citas_por_especialidad.png"
)

plt.close()


# ==========================================
# RANKING MEDICOS
# ==========================================

df_ranking_medicos = (
    df_citas_medico
    .sort_values(
        by="TotalCitas",
        ascending=False
    )
)

# ==========================================
# RANKING ESPECIALIDADES
# ==========================================

df_ranking_especialidades = (
    df_especialidad
    .sort_values(
        by="TotalCitas",
        ascending=False
    )
)

# ==========================================
# DASHBOARD EJECUTIVO V2
# ==========================================

with pd.ExcelWriter(
    "Reportes/Dashboard_Ejecutivo_V2.xlsx",
    engine="openpyxl"
) as writer:

    df_resumen.to_excel(
        writer,
        sheet_name="Resumen",
        index=False
    )

    df_pacientes_ciudad.to_excel(
        writer,
        sheet_name="PacientesCiudad",
        index=False
    )

    df_citas_medico.to_excel(
        writer,
        sheet_name="CitasMedico",
        index=False
    )

    df_especialidad.to_excel(
        writer,
        sheet_name="Especialidades",
        index=False
    )

    df_kpi.to_excel(
        writer,
        sheet_name="KPIs",
        index=False
    )

    df_ranking_medicos.to_excel(
        writer,
        sheet_name="RankingMedicos",
        index=False
    )

    df_ranking_especialidades.to_excel(
        writer,
        sheet_name="RankingEspecialidades",
        index=False
    )



# =====================================================
# ABRIR EXCEL Y FORMATEAR
# =====================================================

wb = load_workbook(
    "Reportes/Dashboard_Ejecutivo_V2.xlsx"
)


# =====================================================
# PORTADA
# =====================================================

portada = wb.create_sheet(
    "Portada",
    0
)

portada.merge_cells("B2:F4")

portada["B2"] = "DASHBOARD EJECUTIVO IPS"

portada["B7"] = "Proyecto SQL Server + Python + Pandas"

portada["B8"] = "Autor: Jairo Fernando Vargas"

portada["B9"] = (
    f"Fecha: "
    f"{datetime.now():%Y-%m-%d %H:%M}"
)


portada["B2"].font = Font(
    size=18,
    bold=True,
    color="FFFFFF"
)

portada["B2"].fill = PatternFill(
    start_color="1F4E78",
    end_color="1F4E78",
    fill_type="solid"
)

portada["B2"].alignment = Alignment(
    horizontal="center",
    vertical="center"
)

portada["B6"] = (
    f"Fecha Generación: "
    f"{datetime.now():%Y-%m-%d %H:%M}"
)


# =====================================================
# FORMATO KPIs
# =====================================================

ws = wb["KPIs"]

# ==========================================
# HOJA KPI
# ==========================================

ws = wb["KPIs"]

# Encabezados originales

color_titulo = PatternFill(
    start_color="1F4E78",
    end_color="1F4E78",
    fill_type="solid"
)

fuente_titulo = Font(
    color="FFFFFF",
    bold=True
)

for celda in ws[1]:

    celda.fill = color_titulo
    celda.font = fuente_titulo


# ==========================================
# KPI CARDS
# ==========================================

azul = PatternFill(
    start_color="5B9BD5",
    end_color="5B9BD5",
    fill_type="solid"
)

verde = PatternFill(
    start_color="70AD47",
    end_color="70AD47",
    fill_type="solid"
)

naranja = PatternFill(
    start_color="ED7D31",
    end_color="ED7D31",
    fill_type="solid"
)

rojo = PatternFill(
    start_color="C00000",
    end_color="C00000",
    fill_type="solid"
)

tarjetas = [
    ("D2", "PACIENTES", total_pacientes, azul),
    ("F2", "MEDICOS", total_medicos, verde),
    ("H2", "CITAS", total_citas, naranja),
    ("J2", "FACTURACION", facturacion, rojo)
]

for titulo, nombre, valor, color in tarjetas:

    ws[titulo] = nombre

    fila_valor = titulo[0] + "3"

    ws[fila_valor] = valor

    ws[titulo].fill = color

    ws[titulo].font = Font(
        bold=True,
        color="FFFFFF",
        size=12
    )

    ws[fila_valor].font = Font(
        bold=True,
        size=16
    )

    ws[titulo].alignment = Alignment(
        horizontal="center"
    )

    ws[fila_valor].alignment = Alignment(
        horizontal="center"
    )

# Formato moneda

ws["J3"].number_format = '$#,##0'


# ==========================================
# SEMAFORO DE FACTURACION
# ==========================================

if facturacion >= 500000:

    ws["L2"] = "🟢 FACTURACION SALUDABLE"

elif facturacion >= 300000:

    ws["L2"] = "🟡 FACTURACION MEDIA"

else:

    ws["L2"] = "🔴 FACTURACION BAJA"

ws["L2"].font = Font(
    bold=True,
    size=12
)

# ==========================================
# AUTO AJUSTAR TODAS LAS HOJAS
# ==========================================

for hoja in wb.worksheets:

    for columna in hoja.columns:

        max_length = 0

        columna_letra = columna[0].column_letter

        for celda in columna:

            try:

                if len(str(celda.value)) > max_length:
                    max_length = len(str(celda.value))

            except:
                pass

        hoja.column_dimensions[
            columna_letra
        ].width = max_length + 3

# =====================================================
# GRAFICOS
# =====================================================

graficos = wb.create_sheet(
    "Graficos"
)

img1 = Image(
    "Reportes/pacientes_por_ciudad.png"
)

graficos.add_image(
    img1,
    "A1"
)

img2 = Image(
    "Reportes/citas_por_especialidad.png"
)

graficos.add_image(
    img2,
    "A20"
)


# =====================================================
# GUARDAR
# =====================================================

wb.save(
    "Reportes/Dashboard_Ejecutivo_V2.xlsx"
)
conexion.close()

print("✅ Conexión cerrada")
print("✅ Dashboard Ejecutivo generado correctamente")
print("✅ Gráficos generados")
print("✅ KPIs generados")